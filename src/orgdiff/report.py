"""Выгрузка результата: Excel (для работы) и HTML (для чтения и демо)."""

from __future__ import annotations

import html
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import views as V
from .pipeline import Result


def sheets(res: Result) -> list[tuple[str, str, list[dict]]]:
    """(название, пояснение, строки) — один источник для Excel, HTML и интерфейса."""
    return [
        ("Подразделения", "Какие подразделения созданы, сохранены или упразднены", V.unit_rows(res)),
        ("Руководители и должности", "Изменения в подчинённости и составе должностей", V.position_rows(res)),
        ("Передача функций", "Какие функции перешли от одного владельца к другому", V.moved_rows(res)),
        ("Потери и дубли", "Функции, которые потеряны или закреплены за несколькими владельцами",
         V.issue_rows(V.by_type(res, "function_lost", "function_lost_candidate", "duplicate"))),
        ("Конфликты интересов", "Сочетания функций, где возможен конфликт интересов",
         V.issue_rows(V.by_type(res, "conflict"))),
        ("Все функции", "Сопоставление каждой функции «было → стало»", V.function_rows(res)),
        ("Изменения пунктов", "Что поменялось в тексте, по пунктам", V.change_rows(res)),
    ]


def write_xlsx(res: Result, path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Вывод"
    ws.append([f"{V.APP_NAME}: {V.APP_TAGLINE}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Было: {res.a.name}   ·   Стало: {res.b.name}"])
    ws.append([V.DISCLAIMER])
    for title, items in V.conclusion(res):
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12, color="305496")
        for it in items:
            ws.append([f"      – {it[2:]}" if it.startswith("— ") else f"•  {it}"])
    for w in res.warnings:
        ws.append([f"⚠ {w}"])
    ws.column_dimensions["A"].width = 150
    for row in ws.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    for title, _, rows in sheets(res):
        _table(wb.create_sheet(title[:31]), V.public(rows))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


WIDTH = {"Функция": 60, "Какие функции": 90, "Что найдено": 55, "Цитаты": 60, "Было": 45, "Стало": 45,
         "Пояснение": 50, "Что изменилось": 45, "Полное название": 55}


def _table(ws, rows: list[dict]) -> None:
    if not rows:
        ws.append(["Нет данных"])
        return
    head = list(rows[0])
    ws.append(head)
    for r in rows:
        ws.append([r[h] for h in head])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, h in enumerate(head, 1):
        longest = max(len(str(r[h]).split("\n")[0]) for r in rows)
        ws.column_dimensions[ws.cell(1, i).column_letter].width = WIDTH.get(h, min(max(len(h), longest) + 2, 34))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_html(res: Result, path: str | Path) -> Path:
    e = html.escape
    warns = "".join(f"<li>{e(w)}</li>" for w in res.warnings)
    model = f"модель {res.model}, запросов к LLM: {res.llm_calls}" if res.llm_used else "без LLM"
    parts = [f"<h2>{e(t)}</h2><p class=muted>{e(d)}</p>{V.table_html(rows)}" for t, d, rows in sheets(res)
             if t != "Все функции"]
    parts.append("<h2>Изменившиеся функции</h2><p class=muted>Сопоставление «было → стало», "
                 "неизменные функции скрыты</p>" + V.table_html(V.function_rows(res, include_unchanged=False)))
    doc = f"""<!doctype html><html lang=ru><head><meta charset=utf-8>
<meta name=viewport content="width=device-width,initial-scale=1">
<title>{V.APP_NAME} — отчёт</title>
<style>{V.CSS}
body{{background:var(--bg);margin:0 auto;padding:24px 16px;max-width:1200px}}
h1{{font-size:24px;margin:0}} h2{{font-size:18px;margin:36px 0 0;border-bottom:2px solid var(--accent);padding-bottom:4px}}
.note{{background:var(--head);padding:10px 14px;border-left:3px solid var(--accent)}}
</style></head><body class=og>
<h1>{V.APP_NAME}</h1><p class=muted>{e(V.APP_TAGLINE)}</p>
<p class=muted>Было: {e(res.a.name)} · Стало: {e(res.b.name)} · {e(model)}</p>
<p class=note>{e(V.DISCLAIMER)}</p>
<h2>Вывод</h2>{V.conclusion_html(res)}
{'<h3>Предупреждения</h3><ul>' + warns + '</ul>' if warns else ''}
{''.join(parts)}
</body></html>"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(doc, encoding="utf-8")
    return path
