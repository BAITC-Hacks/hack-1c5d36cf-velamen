"""Загрузка Word/PDF/Excel/текста в список абзацев."""

from __future__ import annotations

import re
from pathlib import Path

ITEM_START = re.compile(r"^(\d{1,2}(\.\d{1,2})*\.?\s|[а-яa-z]\.\s|[–\-•]\s)")


def load_lines(path: str | Path) -> list[str]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".docx":
        return _docx(path)
    if ext == ".pdf":
        return _pdf(path)
    if ext in (".xlsx", ".xlsm"):
        return _xlsx(path)
    if ext in (".txt", ".md"):
        return [ln.strip() for ln in path.read_text(encoding="utf-8").splitlines() if ln.strip()]
    raise ValueError(f"Неподдерживаемый формат: {ext}")


def _docx(path: Path) -> list[str]:
    import docx

    d = docx.Document(str(path))
    lines = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return lines


def _pdf(path: Path) -> list[str]:
    from pypdf import PdfReader

    raw = []
    for page in PdfReader(str(path)).pages:
        raw.extend(ln.strip() for ln in (page.extract_text() or "").splitlines() if ln.strip())
    # PDF рвёт абзацы по строкам: склеиваем, пока не начнётся новый пункт.
    lines: list[str] = []
    for ln in raw:
        if lines and not ITEM_START.match(ln):
            lines[-1] += " " + ln
        else:
            lines.append(ln)
    return lines


def _xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    lines = []
    for ws in load_workbook(str(path), read_only=True, data_only=True).worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return lines
