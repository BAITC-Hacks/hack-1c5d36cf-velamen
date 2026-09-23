from io import BytesIO

from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter

from orgdiff.chunking import make_chunks
from orgdiff.ingest import parse_file
from orgdiff.models import Document, InputFile, SourceBlock


def save_bytes(document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_numbered_owner_replaces_previous_unnumbered_heading():
    word = WordDocument()
    word.add_paragraph("5. Права и обязанности", style="Heading 1")
    word.add_paragraph("Главный аудитор:", style="Heading 2")
    word.add_paragraph("5.1. Организует работу:")
    word.add_paragraph("5.1.1. утверждает план;")
    word.add_paragraph("5.3. Директор направления внутреннего аудита:")
    word.add_paragraph("5.3.1. готовит предложения;")
    parsed = parse_file(InputFile(name="owners.docx", content=save_bytes(word)))
    assert "Главный аудитор:" in parsed.blocks[3].heading_path
    assert "Главный аудитор:" not in parsed.blocks[-1].heading_path
    assert parsed.blocks[-2].text in parsed.blocks[-1].heading_path


def test_chunks_preserve_every_character_and_parent_context():
    heading = SourceBlock(id="d:h", document_id="d", filename="d.docx", locator="p1", text="Директор не вправе:")
    body = SourceBlock(id="d:b", document_id="d", filename="d.docx", locator="p2", text="абв " * 50,
                       context_ids=[heading.id], heading_path=[heading.text])
    doc = Document(id="d", filename="d.docx", sha256="d", blocks=[heading, body])
    chunks = make_chunks(doc, body_chars=50, context_chars=50)
    recovered = []
    for chunk in chunks:
        if body.id in chunk.body_ids:
            assert heading.id in chunk.body_ids + chunk.context_ids
            recovered.append(chunk.text.split(f"[BODY {body.id}]\n", 1)[1])
    assert "".join(recovered) == body.text


def test_numbered_sibling_ends_previous_roster_context():
    word = WordDocument()
    word.add_paragraph("3. Структура", style="Heading 1")
    word.add_paragraph("3.4. БВА состоит из подразделений:")
    word.add_paragraph("а. Отдел аудита.")
    word.add_paragraph("3.5. Главному аудитору подчиняются:")
    word.add_paragraph("а. Директор.")
    parsed = parse_file(InputFile(name="roster.docx", content=save_bytes(word)))
    assert parsed.blocks[1].text in parsed.blocks[2].heading_path
    assert parsed.blocks[1].text not in parsed.blocks[-1].heading_path


def test_docx_table_cells_keep_order_and_inline_clauses():
    word = WordDocument()
    word.add_paragraph("3.9. Первое условие. 3.10. Второе условие.")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Отдел"
    table.cell(0, 1).text = "Обязанность"
    word.add_paragraph("Следующий абзац")
    parsed = parse_file(InputFile(name="table.docx", content=save_bytes(word)))
    assert [block.locator for block in parsed.blocks] == ["p1", "t1/r1c1/p1", "t1/r1c2/p1", "p2"]
    assert "3.10. Второе условие." in parsed.blocks[0].text


def test_xlsx_missing_formula_value_is_visible_and_incomplete():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Подразделение", "Итого"])
    sheet.append(["Аудит", "=1+2"])
    parsed = parse_file(InputFile(name="formula.xlsx", content=save_bytes(workbook)))
    assert not parsed.read_complete
    assert "=1+2" in parsed.blocks[-1].text
    assert parsed.blocks[0].id in parsed.blocks[-1].context_ids


def test_pdf_without_text_is_incomplete():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    stream = BytesIO()
    writer.write(stream)
    parsed = parse_file(InputFile(name="scan.pdf", content=stream.getvalue()))
    assert not parsed.read_complete
    assert any("Страница 1" in warning for warning in parsed.warnings)
