import asyncio

import pytest

from orgdiff.extract import entity_key, extraction_cache_key, quote_exists, validate_batch
from orgdiff.models import Chunk, Evidence, ExtractionBatch, FunctionFact, SourceBlock


def test_quotes_preserve_negation():
    assert quote_exists("не имеет права", "не  имеет\nправа")
    assert not quote_exists("обязан утверждать", "не имеет права утверждать")
    assert not quote_exists("", "текст")


def test_same_role_in_different_units_is_not_one_entity():
    assert entity_key("Директор проектов", "role", "unit-a") != entity_key("Директор проектов", "role", "unit-b")


def test_cache_includes_parent_heading():
    body = SourceBlock(id="b", document_id="d", filename="a", locator="p2", text="утверждает план")
    heading = SourceBlock(id="h", document_id="d", filename="a", locator="p1", text="Директор:")
    chunk = Chunk(id="c", document_id="d", body_ids=["b"], context_ids=["h"], text=body.text)
    key = extraction_cache_key(chunk, {"b": body, "h": heading}, "model", "1")
    heading.text = "Директор не вправе:"
    assert key != extraction_cache_key(chunk, {"b": body, "h": heading}, "model", "1")


def test_batch_rejects_unprocessed_blocks_and_invented_holders():
    block = SourceBlock(id="b", document_id="d", filename="a", locator="p1", text="утверждает план")
    chunk = Chunk(id="c", document_id="d", body_ids=["b"], text=block.text)
    with pytest.raises(ValueError):
        validate_batch(ExtractionBatch(), chunk, {"b": block})
    fact = FunctionFact(id="f", holder_ids=["ghost"], action="утверждает", object="план",
                        description=block.text, modality="duty", evidence=[Evidence(block_id="b", quote=block.text)])
    with pytest.raises(ValueError):
        validate_batch(ExtractionBatch(functions=[fact], processed_body_ids=["b"]), chunk, {"b": block})


def test_incomplete_extraction_cannot_establish_absence():
    from orgdiff.compare import absence_label
    assert absence_label(coverage_complete=False, found_elsewhere=False) == "unresolved"
    assert absence_label(coverage_complete=True, found_elsewhere=True) == "matched"
    assert absence_label(coverage_complete=True, found_elsewhere=False) == "not_found"


def test_embedding_order_cache_and_validation(tmp_path):
    from types import SimpleNamespace
    from orgdiff.config import Settings
    from orgdiff.llm import Runtime, use_runtime
    from orgdiff.retrieve import embed_texts
    calls = []

    async def create(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(data=[SimpleNamespace(index=1, embedding=[0., 1.]),
                                     SimpleNamespace(index=0, embedding=[1., 0.])], usage=None)

    async def check():
        run = Runtime(Settings(artifact_dir=tmp_path), SimpleNamespace(embeddings=SimpleNamespace(create=create)))
        with use_runtime(run):
            first = await embed_texts(["один", "два"])
            second = await embed_texts(["два", "один"])
            assert first.tolist() == [[1., 0.], [0., 1.]]
            assert second.tolist() == [[0., 1.], [1., 0.]]
            assert len(calls) == 1
            with pytest.raises(ValueError):
                await embed_texts([""])
    asyncio.run(check())


def test_raw_windows_keep_source_and_overlap():
    from orgdiff.retrieve import raw_windows, encoding
    block = SourceBlock(id="b", document_id="d", filename="a", locator="p1", text="Документ русский текст. " * 800)
    windows = raw_windows({block.id: block})
    assert len(windows) > 1
    assert all(w.block_id == "b" and quote_exists(w.text, block.text) for w in windows)
    assert all(len(encoding().encode(w.text)) <= 512 for w in windows)


def test_conflicting_proposals_are_unresolved():
    from orgdiff.compare import reconcile_matches
    from orgdiff.models import FunctionMatch
    matches = [FunctionMatch(before_ids=["a"], after_ids=["b"], changes=["unchanged"], explanation_ru="Совпадает"),
               FunctionMatch(before_ids=["a"], after_ids=["b"], changes=["modality_changed"], explanation_ru="Различается")]
    assert reconcile_matches(matches)[0].changes == ["unresolved"]


def test_adapter_corrects_invalid_response_then_caches(tmp_path):
    from types import SimpleNamespace
    from orgdiff.config import Settings
    from orgdiff.llm import DoctorOutput, Runtime, generate, use_runtime
    calls = []

    async def parse(**kwargs):
        calls.append(kwargs)
        parsed = DoctorOutput(ok=len(calls) > 1)
        response = SimpleNamespace(status="completed", output=[], output_parsed=parsed, usage=None,
                                   model_dump=lambda **kw: {"id": "mock-response"})
        return SimpleNamespace(json=lambda: {"status": "completed"}, parse=lambda: response)

    def validate(result):
        if not result.ok:
            raise ValueError("Неверное значение")

    async def check():
        client = SimpleNamespace(responses=SimpleNamespace(with_raw_response=SimpleNamespace(parse=parse)))
        with use_runtime(Runtime(Settings(artifact_dir=tmp_path), client)):
            for _ in range(2):
                assert (await generate("doctor", {}, DoctorOutput, validate=validate)).ok
        assert len(calls) == 2
        assert calls[0]["store"] is False and calls[0]["service_tier"] == "default"
        assert calls[0]["model"] == "gpt-6-luna" and calls[0]["reasoning"] == {"effort": "none"}
        assert "validation_errors" in calls[1]["input"]
    asyncio.run(check())


def test_failed_extract_is_not_successful_empty_snapshot(monkeypatch, tmp_path):
    from orgdiff import extract
    from orgdiff.config import Settings
    from orgdiff.llm import ModelError, Runtime, use_runtime
    from orgdiff.models import Document

    async def fail(*args, **kwargs):
        raise ModelError("Провайдер недоступен")
    monkeypatch.setattr(extract, "generate", fail)
    block = SourceBlock(id="b", document_id="d", filename="a", locator="p1", text="Текст документа")
    async def check():
        with use_runtime(Runtime(Settings(artifact_dir=tmp_path))):
            snapshot = await extract.extract_snapshot("before", [Document(id="d", filename="a", sha256="d", blocks=[block])])
            assert snapshot.failed_chunk_ids and not snapshot.coverage_complete
            assert snapshot.warnings
    asyncio.run(check())


def test_checkpoint_outputs_and_same_document_match(monkeypatch, tmp_path):
    import json
    from orgdiff import extract
    from orgdiff.config import Settings
    from orgdiff.models import Entity, InputFile
    from orgdiff.pipeline import analyze
    from docx import Document
    from io import BytesIO

    async def extract_fake(task, payload, output_type, **kwargs):
        assert task == "extract"
        block_id = payload["body_ids"][0]
        evidence = [Evidence(block_id=block_id, quote="Директор утверждает план.")]
        batch = ExtractionBatch(entities=[Entity(id="r", name="Директор", kind="role", evidence=evidence)],
            functions=[FunctionFact(id="f", holder_ids=["r"], action="утверждает", object="план", description="Утверждает план",
                modality="duty", owner_resolution="resolved", evidence=evidence)], processed_body_ids=payload["body_ids"])
        kwargs["validate"](batch)
        return batch

    monkeypatch.setattr(extract, "generate", extract_fake)
    from orgdiff import compare
    async def no_calls(*args, **kwargs):
        raise AssertionError("Identical sourced facts should not call the API")
    monkeypatch.setattr(compare, "rank_candidates", no_calls)
    document = Document()
    document.add_paragraph("Директор утверждает план.")
    stream = BytesIO()
    document.save(stream)
    file = InputFile(name="input.docx", content=stream.getvalue())
    settings = Settings(openai_api_key="test-not-a-real-key", artifact_dir=tmp_path)
    async def check():
        await analyze([file, file], [file], tmp_path / "extract", until="extract", settings=settings)
        run = json.loads((tmp_path / "extract/run.json").read_text(encoding="utf-8"))
        assert run["stopped_after"] == "extract" and not run["pipeline_completed"]
        assert "test-not-a-real-key" not in json.dumps(run)
        assert not (tmp_path / "extract/comparison.json").exists()
        result = await analyze([file], [file], tmp_path / "compare", until="compare", settings=settings)
        assert len(result.before.documents) == 1
        assert all(m.changes == ["unchanged"] for m in result.function_matches)
        assert not result.findings and result.stages["risks"] == "not_run"
        assert all((tmp_path / "compare" / name).exists() for name in
                   ["before.snapshot.json", "after.snapshot.json", "sources.json", "comparison.json", "run.json", "usage.json"])
    asyncio.run(check())


def test_same_named_child_units_keep_parents_and_roles():
    from orgdiff.extract import merge_batches
    from orgdiff.models import Entity, Relation, Snapshot
    entities = [Entity(id="p1", name="Аудит", kind="unit"), Entity(id="p2", name="Проекты", kind="unit"),
                Entity(id="u1", name="Отдел контроля", kind="unit"), Entity(id="u2", name="Отдел контроля", kind="unit"),
                Entity(id="r1", name="Директор", kind="role", scope_id="u1"), Entity(id="r2", name="Директор", kind="role", scope_id="u2")]
    relations = [Relation(subject_id="u1", object_id="p1", kind="part_of"), Relation(subject_id="u2", object_id="p2", kind="part_of")]
    snapshot = Snapshot(side="before")
    merge_batches(snapshot, [ExtractionBatch(entities=entities, relations=relations)])
    assert len(snapshot.entities) == 6
    assert len({e.scope_id for e in snapshot.entities if e.kind == "role"}) == 2


def test_comparison_payload_resolves_holder_scope():
    from orgdiff.compare import evidence_payload
    from orgdiff.models import Entity, Snapshot
    role = Entity(id="r", name="Директор", kind="role", scope_id="u")
    unit = Entity(id="u", name="Отдел аудита", kind="unit")
    fact = FunctionFact(id="f", holder_ids=["r"], action="утверждает", object="план", description="План", modality="duty")
    payload = evidence_payload(fact, Snapshot(side="before", entities=[role, unit]))
    assert any(e["name"] == "Отдел аудита" for e in payload["related_entities"])


def test_truncated_json_checked_before_sdk_parses_it(tmp_path):
    from types import SimpleNamespace
    from orgdiff.config import Settings
    from orgdiff.llm import DoctorOutput, Runtime, TruncatedOutput, generate, use_runtime
    def must_not_parse():
        raise AssertionError("SDK must not parse incomplete JSON")
    async def parse(**kwargs):
        return SimpleNamespace(json=lambda: {"status": "incomplete", "incomplete_details": {"reason": "max_output_tokens"},
                                            "usage": {"input_tokens": 5, "output_tokens": 3}}, parse=must_not_parse)
    async def check():
        client = SimpleNamespace(responses=SimpleNamespace(with_raw_response=SimpleNamespace(parse=parse)))
        run = Runtime(Settings(artifact_dir=tmp_path), client)
        with use_runtime(run), pytest.raises(TruncatedOutput):
            await generate("doctor", {}, DoctorOutput)
        assert run.usage["doctor"]["output_tokens"] == 3
        assert not list(tmp_path.rglob("*.json"))
    asyncio.run(check())


@pytest.mark.parametrize("complete,expected", [(False, "unresolved"), (True, "not_found")])
def test_real_comparison_qualifies_negative_search(complete, expected, tmp_path):
    from orgdiff.compare import compare_snapshots
    from orgdiff.config import Settings
    from orgdiff.llm import Runtime, use_runtime
    from orgdiff.models import Document, Entity, Snapshot
    evidence = [Evidence(block_id="b", quote="Директор утверждает план")]
    block = SourceBlock(id="b", document_id="d", filename="a", locator="5.3.2", text=evidence[0].quote)
    before = Snapshot(side="before", coverage_complete=True,
        documents=[Document(id="d", filename="a", sha256="d", blocks=[block])],
        entities=[Entity(id="r", name="Директор", kind="role", evidence=evidence)],
        functions=[FunctionFact(id="f", holder_ids=["r"], action="утверждает", object="план", description="План",
                               modality="duty", owner_resolution="resolved", evidence=evidence)])
    after = Snapshot(side="after", coverage_complete=complete)
    async def check():
        with use_runtime(Runtime(Settings(artifact_dir=tmp_path))):
            result = await compare_snapshots(before, after)
            assert result.function_matches[0].changes == [expected]
            assert bool(result.findings) == complete
            assert result.stages["risks"] == "not_run"
    asyncio.run(check())
